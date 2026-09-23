# -*- coding: utf-8 -*-
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADERS = [
    'CANTIDAD',
    'LARGO DE VETA',
    'ANCHO',
    'DETALLE',
    'CANTO L',
    'CANTO L',
    'CANTO',
    'CANTO',
]

COLOR_HEADER_BG = '2F4F7F'
COLOR_HEADER_FG = 'FFFFFF'
COLOR_CANTO_L_BG = '92D050'
COLOR_CANTO_BG = 'FFFF00'
COLOR_CANTO_FG = '000000'
COLOR_ROW_ALT = 'EEF2F7'
COLOR_WHITE = 'FFFFFF'

# Índices 1-based de columnas CANTO L (verde) y CANTO (amarillo)
CANTO_L_COLUMNS = (5, 6)
CANTO_COLUMNS = (7, 8)


def solid_fill(color):
    return PatternFill(fill_type='solid', fgColor=color)


def load_payload(json_path):
    with open(json_path, encoding='utf-8-sig') as handle:
        return json.load(handle)


def group_modules(rows):
    modules = []
    current_module = None
    current_pieces = []

    for item in rows:
        row_type = item.get('type')
        if row_type == 'module':
            if current_module is not None:
                modules.append((current_module, current_pieces))
            current_module = item.get('label', '')
            current_pieces = []
        elif row_type == 'piece':
            current_pieces.append(item)

    if current_module is not None:
        modules.append((current_module, current_pieces))

    return modules


def piece_invertida(piece):
    return piece.get('invertida') is True


def display_dimensions(piece):
    """Respeta invertida: largo de veta / ancho como en display_dimensions de main.rb."""
    if piece_invertida(piece):
        return piece.get('ancho', 0), piece.get('largo', 0)
    return piece.get('largo', 0), piece.get('ancho', 0)


def display_cantos(piece):
    """Mismo swap que Store.display_cantos cuando invertida=true."""
    arr = int(piece.get('canto_arr', 0) or 0)
    aba = int(piece.get('canto_aba', 0) or 0)
    izq = int(piece.get('canto_izq', 0) or 0)
    der = int(piece.get('canto_der', 0) or 0)

    if piece_invertida(piece):
        return {'arr': izq, 'aba': der, 'izq': arr, 'der': aba}
    return {'arr': arr, 'aba': aba, 'izq': izq, 'der': der}


def color_initials(color_name):
    """Primera letra de cada palabra: 'Seda Nogal' -> 'SN'."""
    words = [w for w in str(color_name or '').strip().split() if w]
    if not words:
        return ''
    return ''.join(word[0].upper() for word in words)


def format_espesor(espesor):
    """Normaliza espesor para DETALLE (punto → coma decimal)."""
    text = str(espesor or '').strip()
    if not text:
        return ''
    return text.replace('.', ',')


def canto_mark(value):
    """'1' si el lado tiene canto (rojo o azul), vacío si no."""
    return '1' if int(value or 0) != 0 else ''


def canto_detalle_part(piece, canto_value):
    """Parte DETALLE de un lado: 'SN - 0,45' o None si sin canto."""
    value = int(canto_value or 0)
    if value == 0:
        return None

    if value == 1:
        color = piece.get('canto_rojo_color', '')
        espesor = piece.get('canto_rojo_espesor', '')
    else:
        color = piece.get('canto_azul_color', '')
        espesor = piece.get('canto_azul_espesor', '')

    initials = color_initials(color)
    esp = format_espesor(espesor)
    if not initials and not esp:
        return None
    if initials and esp:
        return f'{initials} - {esp}'
    return initials or esp


def build_detalle(piece, cantos):
    """
    DETALLE único por color/espesor distinto entre lados marcados.
    Varios: unidos con ' / ' (ej. 'SN - 0,45 / B - 2').
    """
    parts = []
    seen = set()
    # Mismo orden de lados que CorteCloud: L1=aba, L2=arr, A1=izq, A2=der
    for side in ('aba', 'arr', 'izq', 'der'):
        part = canto_detalle_part(piece, cantos.get(side, 0))
        if part and part not in seen:
            seen.add(part)
            parts.append(part)
    return ' / '.join(parts)


def write_header_row(sheet, row_index):
    base_fill = solid_fill(COLOR_HEADER_BG)
    base_font = Font(bold=True, color=COLOR_HEADER_FG)
    canto_l_fill = solid_fill(COLOR_CANTO_L_BG)
    canto_fill = solid_fill(COLOR_CANTO_BG)
    canto_font = Font(bold=True, color=COLOR_CANTO_FG)

    for column_index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=row_index, column=column_index, value=header)
        cell.alignment = Alignment(horizontal='center')
        if column_index in CANTO_L_COLUMNS:
            cell.fill = canto_l_fill
            cell.font = canto_font
        elif column_index in CANTO_COLUMNS:
            cell.fill = canto_fill
            cell.font = canto_font
        else:
            cell.fill = base_fill
            cell.font = base_font


def write_piece_row(sheet, row_index, item, use_alt_fill):
    largo, ancho = display_dimensions(item)
    cantos = display_cantos(item)

    # CANTO L = lados del largo de veta (aba, arr); CANTO = lados del ancho (izq, der)
    values = [
        item.get('cantidad', 0),
        largo,
        ancho,
        build_detalle(item, cantos),
        canto_mark(cantos['aba']),
        canto_mark(cantos['arr']),
        canto_mark(cantos['izq']),
        canto_mark(cantos['der']),
    ]

    fill_color = COLOR_ROW_ALT if use_alt_fill else COLOR_WHITE
    fill = solid_fill(fill_color)

    for column_index, value in enumerate(values, start=1):
        cell = sheet.cell(row=row_index, column=column_index, value=value)
        cell.fill = fill
        if column_index in (5, 6, 7, 8):
            cell.alignment = Alignment(horizontal='center')


def adjust_column_widths(sheet):
    widths = {
        1: 12,
        2: 16,
        3: 12,
        4: 22,
        5: 10,
        6: 10,
        7: 10,
        8: 10,
    }
    for column_index, width in widths.items():
        column_letter = sheet.cell(row=1, column=column_index).column_letter
        sheet.column_dimensions[column_letter].width = width


def write_xlsx(output_path, payload):
    modules = group_modules(payload.get('rows', []))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'MaderAmericana'

    write_header_row(sheet, 1)
    row_index = 2
    use_alt_fill = False

    for _module_label, pieces in modules:
        for piece in pieces:
            write_piece_row(sheet, row_index, piece, use_alt_fill)
            use_alt_fill = not use_alt_fill
            row_index += 1

    adjust_column_widths(sheet)
    workbook.save(output_path)


def main():
    if len(sys.argv) < 3:
        sys.stderr.write('Uso: export_maderamerica.py salida.xlsx datos.json\n')
        sys.exit(1)

    output_path = sys.argv[1]
    json_path = sys.argv[2]
    payload = load_payload(json_path)
    write_xlsx(output_path, payload)


if __name__ == '__main__':
    main()
